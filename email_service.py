import os
import io
import csv
import re
import base64
import requests
from typing import List, Dict, Any, Optional, Tuple
from dotenv import load_dotenv
import openpyxl

load_dotenv()

BREVO_API_KEY = os.getenv("BREVO_API_KEY")
SENDER_EMAIL = os.getenv("SENDER_EMAIL", "gamerkamal028@gmail.com")
BREVO_API_URL = "https://api.brevo.com/v3/smtp/email"


def build_html_wrapper(subject: str, body_html: str) -> str:
    """Wraps rich HTML body content in an Apple-inspired email frame."""
    return f"""<!DOCTYPE html>
<html>
    <body style="margin: 0; padding: 40px 20px; background-color: #f5f5f7; font-family: -apple-system, BlinkMacSystemFont, 'SF Pro Text', 'Segoe UI', Roboto, Helvetica, Arial, sans-serif; color: #1d1d1f;">
        <div style="max-width: 620px; margin: 0 auto; background-color: #ffffff; border-radius: 18px; border: 1px solid rgba(0, 0, 0, 0.08); padding: 44px; box-sizing: border-box; box-shadow: 0 4px 20px rgba(0,0,0,0.03);">
            <div style="font-size: 11px; font-weight: 700; color: #86868b; text-transform: uppercase; letter-spacing: 0.8px; margin-bottom: 14px;">Mail Dispatch Studio</div>
            <h1 style="font-size: 24px; font-weight: 600; line-height: 1.3; margin: 0 0 24px 0; color: #1d1d1f; letter-spacing: -0.3px;">{subject}</h1>
            <div style="font-size: 15px; line-height: 1.65; color: #333336;">
                {body_html}
            </div>
            <hr style="border: none; border-top: 1px solid #f0f0f2; margin: 36px 0 20px 0;">
            <p style="font-size: 12px; color: #86868b; margin: 0; line-height: 1.4;">
                Dispatched securely via Mail Dispatch Studio &bull; Designed with Apple Aesthetics
            </p>
        </div>
    </body>
</html>"""


def parse_recipients_file(filename: str, content: bytes) -> List[Dict[str, str]]:
    """Parses .csv or .xlsx spreadsheets into normalized row dictionaries."""
    records: List[Dict[str, str]] = []
    ext = filename.lower().split(".")[-1]

    if content.startswith(b"PK\x03\x04"):
        ext = "xlsx"

    if ext == "csv":
        raw_text = content.decode("utf-8-sig", errors="ignore")
        normalized_text = (
            raw_text.replace("\r\r\n", "\n")
            .replace("\r\n", "\n")
            .replace("\r", "\n")
        )

        sample = normalized_text[:2048]
        delimiter = ","
        if sample.count(";") > sample.count(",") and sample.count(";") > sample.count("\t"):
            delimiter = ";"
        elif sample.count("\t") > sample.count(","):
            delimiter = "\t"

        text_stream = io.StringIO(normalized_text, newline="")
        reader = csv.DictReader(text_stream, delimiter=delimiter)

        for row in reader:
            if not row:
                continue
            clean_row = {
                str(k).strip(): str(v).strip()
                for k, v in row.items()
                if k is not None and v is not None
            }
            if clean_row:
                records.append(clean_row)

    elif ext in ["xlsx", "xls"]:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if rows:
            headers = [
                str(h).strip() if h is not None else f"col_{idx}"
                for idx, h in enumerate(rows[0])
            ]
            for data_row in rows[1:]:
                row_dict = {}
                for idx, cell in enumerate(data_row):
                    if idx < len(headers):
                        row_dict[headers[idx]] = str(cell).strip() if cell is not None else ""
                if any(row_dict.values()):
                    records.append(row_dict)

    valid_records = []
    for rec in records:
        email_key = None
        for k in rec.keys():
            if k.lower() in ["email", "e-mail", "mail", "recipient", "to", "email address", "email_id"]:
                email_key = k
                break

        if not email_key:
            for k, val in rec.items():
                if "@" in val and "." in val:
                    email_key = k
                    break

        if email_key and rec.get(email_key):
            rec["__email__"] = rec[email_key].strip()
            valid_records.append(rec)

    return valid_records


def substitute_placeholders(template: str, context: Dict[str, str]) -> str:
    """Replaces placeholders like {name} or {company} case-insensitively."""
    lookup = {k.lower().strip(): v for k, v in context.items() if not k.startswith("__")}

    def replacer(match: re.Match) -> str:
        key = match.group(1).lower().strip()
        return lookup.get(key, match.group(0))

    return re.sub(r"\{([a-zA-Z0-9_]+)\}", replacer, template)


def send_brevo_request(
    to_list: List[str],
    subject: str,
    body: str,
    cc_list: Optional[List[str]] = None,
    bcc_list: Optional[List[str]] = None,
    attachment_name: Optional[str] = None,
    attachment_data: Optional[bytes] = None,
    scheduled_at: Optional[str] = None
) -> Tuple[bool, str]:
    """Transmits an email payload to Brevo with rich HTML and plain text fallback."""
    if not BREVO_API_KEY:
        return False, "BREVO_API_KEY is not configured."

    # Strip HTML tags for clean plain-text fallback
    plain_text_fallback = re.sub(r"<[^>]+>", " ", body)
    plain_text_fallback = re.sub(r"\s+", " ", plain_text_fallback).strip()

    payload: Dict[str, Any] = {
        "sender": {"name": "Mail Dispatch Studio", "email": SENDER_EMAIL},
        "to": [{"email": e} for e in to_list],
        "subject": subject,
        "htmlContent": build_html_wrapper(subject, body),
        "textContent": plain_text_fallback
    }

    if scheduled_at and scheduled_at.strip():
        payload["scheduledAt"] = scheduled_at.strip()

    if cc_list:
        payload["cc"] = [{"email": e} for e in cc_list]
    if bcc_list:
        payload["bcc"] = [{"email": e} for e in bcc_list]

    if attachment_name and attachment_data:
        payload["attachment"] = [
            {
                "name": attachment_name,
                "content": base64.b64encode(attachment_data).decode("utf-8")
            }
        ]

    headers = {
        "accept": "application/json",
        "api-key": BREVO_API_KEY.strip(),
        "content-type": "application/json"
    }

    try:
        res = requests.post(BREVO_API_URL, json=payload, headers=headers, timeout=15)
        if res.status_code in [200, 201, 202]:
            return True, res.json().get("messageId", "sent")
        else:
            return False, f"Brevo API error ({res.status_code}): {res.text}"
    except Exception as e:
        return False, str(e)


def send_bulk_personalized(
    records: List[Dict[str, str]],
    subject_template: str,
    body_template: str,
    cc_list: Optional[List[str]] = None,
    bcc_list: Optional[List[str]] = None,
    attachment_name: Optional[str] = None,
    attachment_data: Optional[bytes] = None,
    scheduled_at: Optional[str] = None
) -> Dict[str, Any]:
    """Sends personalized copies to each recipient parsed from a spreadsheet."""
    sent_count = 0
    failures = []

    for row in records:
        target_email = row["__email__"]
        personalized_subject = substitute_placeholders(subject_template, row)
        personalized_body = substitute_placeholders(body_template, row)

        ok, msg = send_brevo_request(
            to_list=[target_email],
            subject=personalized_subject,
            body=personalized_body,
            cc_list=cc_list,
            bcc_list=bcc_list,
            attachment_name=attachment_name,
            attachment_data=attachment_data,
            scheduled_at=scheduled_at
        )

        if ok:
            sent_count += 1
        else:
            failures.append({"email": target_email, "error": msg})

    return {
        "total": len(records),
        "successful": sent_count,
        "failed": len(failures),
        "failure_details": failures,
        "scheduled_at": scheduled_at
    }